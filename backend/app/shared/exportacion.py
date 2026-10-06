"""Exportación de tablas a Excel, PDF y HTML (bitácora y reportes personalizados)."""

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from html import escape
from io import BytesIO
from typing import Any

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ZONA_BOLIVIA = timezone(timedelta(hours=-4))

# Las fuentes base del PDF solo tienen Latin-1: estos caracteres se cambian por equivalentes.
_EQUIVALENTES_LATIN1 = str.maketrans({"—": "-", "–": "-", "•": "*", "“": '"', "”": '"', "‘": "'", "’": "'", "…": "..."})


@dataclass
class Tabla:
    titulo: str
    encabezados: list[str]
    filas: list[list[Any]]
    # Líneas bajo el título: quién lo generó, cuándo y con qué filtros.
    descripcion: list[str] = field(default_factory=list)
    # 'numero' alinea a la derecha; cualquier otro valor, a la izquierda.
    tipos: list[str] | None = None


def hora_bolivia(valor: datetime) -> datetime:
    return valor.astimezone(ZONA_BOLIVIA) if valor.tzinfo else valor


def texto_celda(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, datetime):
        return hora_bolivia(valor).strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal):
        return f"{valor:,.2f}"
    return str(valor)


def _es_numero(tabla: Tabla, indice: int) -> bool:
    return bool(tabla.tipos) and tabla.tipos[indice] == "numero"


def a_excel(tabla: Tabla) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "".join(c for c in tabla.titulo if c not in "[]:*?/\\")[:31] or "Reporte"
    ws.append([tabla.titulo])
    ws["A1"].font = Font(bold=True, size=14)
    for linea in tabla.descripcion:
        ws.append([linea])
    ws.append([])
    fila_encabezado = ws.max_row + 1
    ws.append(tabla.encabezados)
    for celda in ws[fila_encabezado]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E3A8A")
        celda.alignment = Alignment(vertical="center")
    for fila in tabla.filas:
        ws.append([_valor_excel(v) for v in fila])
    for indice, encabezado in enumerate(tabla.encabezados, start=1):
        largo = max([len(encabezado)] + [len(texto_celda(f[indice - 1])) for f in tabla.filas[:500]])
        ws.column_dimensions[get_column_letter(indice)].width = min(max(largo + 2, 10), 60)
        if any(isinstance(f[indice - 1], datetime) for f in tabla.filas[:50]):
            for (celda,) in ws.iter_rows(min_row=fila_encabezado + 1, min_col=indice, max_col=indice):
                celda.number_format = "dd/mm/yyyy hh:mm"
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)
    if tabla.filas:
        ws.auto_filter.ref = f"A{fila_encabezado}:{get_column_letter(len(tabla.encabezados))}{ws.max_row}"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _valor_excel(valor: Any) -> Any:
    # Excel no guarda zona horaria: se escribe la hora de Bolivia.
    if isinstance(valor, datetime):
        return hora_bolivia(valor).replace(tzinfo=None)
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def _latin1(texto: str) -> str:
    return texto.translate(_EQUIVALENTES_LATIN1).encode("latin-1", "replace").decode("latin-1")


def a_pdf(tabla: Tabla) -> bytes:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 9, _latin1(tabla.titulo), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.set_text_color(90, 90, 90)
    for linea in tabla.descripcion:
        pdf.multi_cell(0, 4.5, _latin1(linea), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(2)

    # Ancho de cada columna según el largo de su contenido, con tope para los textos largos.
    largos = [
        min(max([len(e)] + [len(texto_celda(f[i])) for f in tabla.filas[:200]]), 45)
        for i, e in enumerate(tabla.encabezados)
    ]
    total = sum(largos) or 1
    anchos = [max(l / total * pdf.epw, 14) for l in largos]
    pdf.set_font("Helvetica", "", 7.5)
    with pdf.table(
        col_widths=anchos,
        width=sum(anchos) if sum(anchos) <= pdf.epw else pdf.epw,
        line_height=4.2,
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(30, 58, 138)),
        cell_fill_color=(241, 245, 249),
        cell_fill_mode="ROWS",
        first_row_as_headings=True,
    ) as tabla_pdf:
        encabezado = tabla_pdf.row()
        for texto in tabla.encabezados:
            encabezado.cell(_latin1(texto))
        for fila in tabla.filas:
            renglon = tabla_pdf.row()
            for indice, valor in enumerate(fila):
                renglon.cell(_latin1(texto_celda(valor))[:300], align="RIGHT" if _es_numero(tabla, indice) else "LEFT")
    return bytes(pdf.output())


def a_html(tabla: Tabla) -> bytes:
    filas = "\n".join(
        "<tr>"
        + "".join(
            f'<td class="{"num" if _es_numero(tabla, i) else ""}">{escape(texto_celda(v))}</td>' for i, v in enumerate(fila)
        )
        + "</tr>"
        for fila in tabla.filas
    )
    encabezados = "".join(f"<th>{escape(e)}</th>" for e in tabla.encabezados)
    descripcion = "".join(f"<p>{escape(l)}</p>" for l in tabla.descripcion)
    vacio = f'<tr><td colspan="{len(tabla.encabezados)}" class="vacio">Sin resultados para los filtros elegidos.</td></tr>'
    documento = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(tabla.titulo)}</title>
<style>
  body {{ font-family: system-ui, -apple-system, 'Segoe UI', Roboto, sans-serif; color: #0f172a; margin: 24px; }}
  h1 {{ font-size: 20px; margin: 0 0 6px; }}
  .meta p {{ margin: 2px 0; color: #475569; font-size: 13px; }}
  .marco {{ overflow-x: auto; margin-top: 16px; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  th {{ background: #1e3a8a; color: #fff; text-align: left; padding: 8px 10px; position: sticky; top: 0; }}
  td {{ padding: 7px 10px; border-bottom: 1px solid #e2e8f0; vertical-align: top; }}
  tr:nth-child(even) td {{ background: #f8fafc; }}
  td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  td.vacio {{ text-align: center; color: #64748b; padding: 24px; }}
  footer {{ margin-top: 16px; color: #94a3b8; font-size: 12px; }}
  @media print {{ body {{ margin: 0; }} th {{ position: static; }} }}
</style>
</head>
<body>
<h1>{escape(tabla.titulo)}</h1>
<div class="meta">{descripcion}</div>
<div class="marco">
<table>
<thead><tr>{encabezados}</tr></thead>
<tbody>
{filas or vacio}
</tbody>
</table>
</div>
<footer>EGRESA · {len(tabla.filas)} filas</footer>
</body>
</html>
"""
    return documento.encode("utf-8")
